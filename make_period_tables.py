#!/usr/bin/env python3
"""Build the KEPROBA period-performance workbook (quarter / semi-annual /
annual) from raw KNBS monthly trade data.

Sheets (same shape as the completed half-year workbook):

    EXP PDT   export products      (KEPROBA product groups + HS chapters)
    EXP DES   export destinations
    IMP PDT   import products      (HS chapters)
    IMP SUP   import supplying countries
    BOT       export / import / balance of trade
    DATA PARTNER / DATA GROUP      monthly aggregates the summary cells
                                   are calculated from

Every derived cell is a live Excel formula with a cached result, so the
workbook recalculates in Excel while the report generator can still read
plain numbers.  The formulas:

    sub-period    = SUMIFS over the DATA sheets (traceable to raw months)
    period total  = SUM of its sub-period cells
    change        = period total (year) - period total (year-1)
    %             = change / period total (year-1)     <- correct base
    all other     = grand total - SUM(listed rows)
    grand total   = SUMIFS over the DATA sheets

Change and % columns are colour coded (green up / red down) for increases
and decreases.  Exports are domestic exports + re-exports at every level
(month, quarter, half year, year).

Usage
-----
    python make_period_tables.py --data-dir KNBS --out-dir output/hy2026 \\
        --period semi --year 2026
"""

import argparse
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

from make_tables import _save_workbook  # noqa: E402
from make_quarterly_tables import FMT_CHANGE, FMT_PCT, FMT_VALUE  # noqa: E402
from period_data import load_monthly, period_spec, data_sheets  # noqa: E402

FONT = "Century Gothic"
HDR_FONT = Font(name=FONT, size=11, bold=True, color="000000")
BOLD = Font(name=FONT, size=10, bold=True)
PLAIN = Font(name=FONT, size=10)
THIN = Side(style="thin", color="B0B0B0")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

# column letters inside the DATA sheets
# DATA PARTNER: Year Month Side Partner Value          -> A B C D E
# DATA GROUP:   Year Month Side Group GroupLabel Value -> A B C D E F
DP, DG = "'DATA PARTNER'", "'DATA GROUP'"


class Book:
    """Workbook with a per-sheet {cell_ref: cached_value} formula cache."""

    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self.cache = {}

    def sheet(self, name):
        ws = self.wb.create_sheet(name[:31])
        self.cache[ws.title] = {}
        return ws

    def formula(self, ws, row, col, formula, cached, numfmt=None,
                bold=False, align=None):
        cell = ws.cell(row, col)
        cell.value = formula
        cell.border = GRID
        cell.font = BOLD if bold else PLAIN
        cell.alignment = align or RIGHT
        if numfmt:
            cell.number_format = numfmt
        self.cache[ws.title]["%s%d" % (get_column_letter(col), row)] = cached

    def plain(self, ws, row, col, value, bold=False, align=None,
              numfmt=None):
        cell = ws.cell(row, col)
        cell.value = value
        cell.border = GRID
        cell.font = BOLD if bold else PLAIN
        cell.alignment = align or RIGHT
        if numfmt:
            cell.number_format = numfmt

    def cached(self, ws, row, col):
        return self.cache[ws.title].get(
            "%s%d" % (get_column_letter(col), row), 0.0) or 0.0

    def save(self, path):
        _save_workbook(self.wb, path, self.cache)


def sumifs(sheet, val_col, year, months, side, key_col=None, key_ref=None):
    """SUMIFS over one DATA sheet for one sub-period / key cell."""
    ms = sorted(months)
    parts = [sheet + "!$%s:$%s" % (val_col, val_col),
             "%s!$A:$A,%d" % (sheet, year),
             "%s!$B:$B,\">=%d\"" % (sheet, ms[0]),
             "%s!$B:$B,\"<=%d\"" % (sheet, ms[-1]),
             "%s!$C:$C,\"%s\"" % (sheet, side)]
    if key_col is not None:
        parts.append("%s!$%s:$%s,%s" % (sheet, key_col, key_col, key_ref))
    return "SUMIFS(%s)" % ",".join(parts)


class PeriodTables:
    def __init__(self, df, spec):
        self.df = df
        self.spec = spec

    # -- cached aggregation matching a SUMIFS -----------------------------
    def value(self, yr, ms, side, key=None, key_kind=None):
        d = self.df[(self.df["Year"] == yr) & (self.df["Side"] == side)
                    & (self.df["Month"].isin(ms))]
        if key is not None:
            col = "GroupLabel" if key_kind == "group" else "Partner"
            d = d[d[col] == key]
        return float(d["Value"].sum())

    def items(self, side, label_col, code_col=None):
        spec = self.spec
        d = self.df[(self.df["Side"] == side) & (self.df["Year"] == spec["year"])
                    & (self.df["Month"].isin(spec["months"]))]
        keys = [label_col] + ([code_col] if code_col else [])
        grp = d.groupby(keys, dropna=False)["Value"].sum().reset_index()
        grp = grp.sort_values("Value", ascending=False)
        out = []
        for i, rec in enumerate(grp.itertuples(index=False), start=1):
            out.append({"rank": i,
                        "label": getattr(rec, label_col),
                        "code": getattr(rec, code_col) if code_col else "",
                        "value": float(rec.Value)})
        return out

    # -- one summary sheet -----------------------------------------------
    def write_rank_sheet(self, book, name, items, label_cols, side,
                         key_kind, other_label, top_n):
        ws = book.sheet(name)
        spec = self.spec
        subs = spec["subs"]
        n = len(subs)
        y0, y1 = spec["prev_year"], spec["year"]
        nlab = len(label_cols)
        # columns: labels | yr-1 subs | yr-1 total | yr-2 subs | yr-2 total
        #          | change | %
        c0 = nlab + 1
        c1 = nlab + n + 2
        t0 = nlab + n + 1
        t1 = nlab + 2 * n + 2
        chg = nlab + 2 * n + 3

        # ---- header ----
        for j, (text, _k) in enumerate(label_cols, start=1):
            book.plain(ws, 1, j, text, bold=True, align=CENTER)
            book.plain(ws, 2, j, "", align=CENTER)
            ws.merge_cells(start_row=1, start_column=j,
                           end_row=2, end_column=j)
        for base, yr in ((c0, y0), (c1, y1)):
            for k, (sub, _m) in enumerate(subs):
                book.plain(ws, 1, base + k, sub, bold=True, align=CENTER)
                book.plain(ws, 2, base + k, "Ksh. Billion", align=CENTER)
            book.plain(ws, 1, base + n, "%s %d" % (spec["span_short"], yr),
                       bold=True, align=CENTER)
            book.plain(ws, 2, base + n, "Ksh. Billion", bold=True, align=CENTER)
        book.plain(ws, 1, chg, "Change in %d-%d" % (y0, y1), bold=True,
                   align=CENTER)
        book.plain(ws, 1, chg + 1, "", bold=True, align=CENTER)
        ws.merge_cells(start_row=1, start_column=chg,
                       end_row=1, end_column=chg + 1)
        book.plain(ws, 2, chg, "Ksh. Billion", align=CENTER)
        book.plain(ws, 2, chg + 1, "%", bold=True, align=CENTER)

        sheet = DG if key_kind == "group" else DP
        vcol = "F" if key_kind == "group" else "E"
        kcol = "E" if key_kind == "group" else "D"

        r = 3
        first = r
        for it in items[:top_n]:
            self._emit(book, ws, it, r, label_cols, side, key_kind,
                       sheet, vcol, kcol, c0, c1, t0, t1, chg, n)
            r += 1
        last = r - 1

        # ---- All other (grand total minus the listed rows) ----
        ao = r
        book.plain(ws, ao, nlab, other_label, bold=True, align=LEFT)
        for base, yr in ((c0, y0), (c1, y1)):
            for k, (_s, ms) in enumerate(subs):
                f = sumifs(sheet, vcol, yr, ms, side)
                grand = self.value(yr, ms, side)
                listed = sum(book.cached(ws, rr, base + k)
                             for rr in range(first, last + 1))
                if last >= first:
                    letter = get_column_letter(base + k)
                    expr = "%s-SUM(%s%d:%s%d)" % (
                        f, letter, first, letter, last)
                else:
                    # no listed rows: the formula would sum through itself
                    # (first = last + 1) and Excel would report a circular
                    # reference, so "All other" is the grand total alone.
                    expr = f
                book.formula(ws, ao, base + k, "=" + expr,
                             grand - listed, numfmt=FMT_VALUE)
            self._total(book, ws, ao, base, base + n, n, c0, c1, t0, t1, chg)
        self._change(book, ws, ao, t0, t1, chg)

        # ---- Grand total ----
        gt = r + 1
        book.plain(ws, gt, nlab, "Grand Total", bold=True, align=LEFT)
        for base, yr in ((c0, y0), (c1, y1)):
            for k, (_s, ms) in enumerate(subs):
                f = sumifs(sheet, vcol, yr, ms, side)
                book.formula(ws, gt, base + k, "=" + f,
                             self.value(yr, ms, side), numfmt=FMT_VALUE,
                             bold=True)
            self._total(book, ws, gt, base, base + n, n, c0, c1, t0, t1, chg,
                        bold=True)
        self._change(book, ws, gt, t0, t1, chg, bold=True)

        ws.freeze_panes = "A3"
        widths = [10, 10, 40] if nlab == 3 else [10, 32]
        for j, w in enumerate(widths[:nlab], start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for j in range(nlab + 1, chg + 2):
            ws.column_dimensions[get_column_letter(j)].width = 13

    def _emit(self, book, ws, it, r, label_cols, side, key_kind,
              sheet, vcol, kcol, c0, c1, t0, t1, chg, n):
        for j, (_t, kind) in enumerate(label_cols, start=1):
            val = it["label"] if kind == "label" else it.get(kind, "")
            book.plain(ws, r, j, val,
                       align=LEFT if kind == "label" else CENTER)
        key_ref = "$%s%d" % (get_column_letter(len(label_cols)), r)
        for base, yr in ((c0, self.spec["prev_year"]), (c1, self.spec["year"])):
            for k, (_s, ms) in enumerate(self.spec["subs"]):
                f = sumifs(sheet, vcol, yr, ms, side, kcol, key_ref)
                book.formula(ws, r, base + k, "=" + f,
                             self.value(yr, ms, side, it["label"], key_kind),
                             numfmt=FMT_VALUE)
            self._total(book, ws, r, base, base + n, n, c0, c1, t0, t1, chg)
        self._change(book, ws, r, t0, t1, chg)

    def _total(self, book, ws, r, base, total_col, n, c0, c1, t0, t1, chg,
               bold=False):
        a = get_column_letter(base)
        b = get_column_letter(base + n - 1)
        cached = sum(book.cached(ws, r, base + k) for k in range(n))
        book.formula(ws, r, total_col,
                     "=SUM(%s%d:%s%d)" % (a, r, b, r), cached,
                     numfmt=FMT_VALUE, bold=bold)

    def _change(self, book, ws, r, t0, t1, chg, bold=False):
        a = "%s%d" % (get_column_letter(t0), r)
        b = "%s%d" % (get_column_letter(t1), r)
        prev = book.cached(ws, r, t0)
        cur = book.cached(ws, r, t1)
        book.formula(ws, r, chg, "=%s-%s" % (b, a), cur - prev,
                     numfmt=FMT_CHANGE, bold=bold)
        pct = (cur - prev) / prev if prev else 0.0
        book.formula(ws, r, chg + 1,
                     '=IF(%s=0,"",(%s-%s)/%s)' % (a, b, a, a), pct,
                     numfmt=FMT_PCT, bold=bold)

    def write_bot(self, book):
        ws = book.sheet("BOT")
        spec = self.spec
        subs = spec["subs"]
        n = len(subs)
        y0, y1 = spec["prev_year"], spec["year"]
        c0, c1 = 2, 2 + n + 1
        book.plain(ws, 3, 1, "Year", bold=True, align=CENTER)
        for base, yr in ((c0, y0), (c1, y1)):
            for k, (sub, _m) in enumerate(subs):
                book.plain(ws, 3, base + k, sub, bold=True, align=CENTER)
            book.plain(ws, 3, base + n, "%s %d" % (spec["span_short"], yr),
                       bold=True, align=CENTER)
        r = 4
        for label, side in (("Export", "Exports"), ("Import", "Imports"),
                            ("Bal of Trade", None)):
            book.plain(ws, r, 1, label, bold=True, align=LEFT)
            for base, yr in ((c0, y0), (c1, y1)):
                for k, (_s, ms) in enumerate(subs):
                    col = get_column_letter(base + k)
                    if side:
                        f = "=" + sumifs(DP, "E", yr, ms, side)
                        book.formula(ws, r, base + k, f,
                                     self.value(yr, ms, side),
                                     numfmt=FMT_VALUE)
                    else:
                        f = "=%s4-%s5" % (col, col)
                        cached = (book.cached(ws, 4, base + k)
                                  - book.cached(ws, 5, base + k))
                        book.formula(ws, r, base + k, f, cached,
                                     numfmt=FMT_VALUE)
                a = get_column_letter(base)
                b = get_column_letter(base + n - 1)
                cached = sum(book.cached(ws, r, base + k) for k in range(n))
                book.formula(ws, r, base + n,
                             "=SUM(%s%d:%s%d)" % (a, r, b, r), cached,
                             numfmt=FMT_VALUE, bold=True)
            r += 1
        ws.column_dimensions["A"].width = 16

    def write_meta(self, book):
        """Period + source metadata consumed by generate_period_report.py."""
        ws = book.sheet("META")
        spec = self.spec
        if "SourceKind" in self.df.columns:
            kinds = sorted(set(self.df["SourceKind"].dropna()))
        else:
            kinds = []
        rows = [("period", spec["period"]),
                ("year", spec["year"]),
                ("prev_year", spec["prev_year"]),
                ("span", spec["span"]),
                ("span_short", spec["span_short"]),
                ("subs", "|".join(s for s, _m in spec["subs"])),
                ("months", "|".join(str(m) for m in spec["months"])),
                ("source", " and ".join(kinds) if kinds else "KNBS")]
        for i, (k, v) in enumerate(rows, start=1):
            book.plain(ws, i, 1, k, bold=True, align=LEFT)
            book.plain(ws, i, 2, v, align=LEFT)
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 40

    def write_data(self, book, sheets):
        for name, df in sheets.items():
            ws = book.sheet(name)
            for j, col in enumerate(df.columns, start=1):
                book.plain(ws, 1, j, col, bold=True, align=CENTER)
            for i, rec in enumerate(df.itertuples(index=False), start=2):
                for j, col in enumerate(df.columns, start=1):
                    book.plain(ws, i, j, getattr(rec, col))
            for j in range(1, len(df.columns) + 1):
                ws.column_dimensions[get_column_letter(j)].width = 15
            ws.freeze_panes = "A2"


def build(data_dir, out_dir, period, year=None, half=1, quarter=1,
          top_products=20, top_markets=25, prev_year=None):
    # validate the period arguments before the slow data load
    period_spec(period, year or 0, half=half, quarter=quarter,
                prev_year=prev_year)
    print("Reading raw monthly data from %s ..." % data_dir)
    df = load_monthly(data_dir)
    if year is None:
        # the newest year in the uploads is the review year (the comparison
        # is always that same sub-period one year earlier)
        year = int(df["Year"].max())
        print("  [info] review year inferred from data: %d" % year)
    spec = period_spec(period, year, half=half, quarter=quarter,
                       prev_year=prev_year)
    print("Review period: %s %d vs %d (%d months)"
          % (spec["span"], spec["year"], spec["prev_year"], len(spec["months"])))

    pt = PeriodTables(df, spec)
    book = Book()
    pt.write_meta(book)
    pt.write_data(book, data_sheets(df))

    yr = spec["year"]
    pt.write_rank_sheet(
        book, "EXP PDT", pt.items("Exports", "GroupLabel", "Group"),
        [("HS code", "code"), ("Rank in %d" % yr, "rank"),
         ("Product Label", "label")],
        "Exports", "group", "All other products", top_products)
    pt.write_rank_sheet(
        book, "EXP DES", pt.items("Exports", "Partner"),
        [("Rank in %d" % yr, "rank"), ("Destination", "label")],
        "Exports", "partner", "All other countries", top_markets)
    pt.write_rank_sheet(
        book, "IMP PDT", pt.items("Imports", "GroupLabel", "Group"),
        [("HS code", "code"), ("Rank in %d" % yr, "rank"),
         ("Product Label", "label")],
        "Imports", "group", "All other products", top_products)
    pt.write_rank_sheet(
        book, "IMP SUP", pt.items("Imports", "Partner"),
        [("Rank in %d" % yr, "rank"), ("Source", "label")],
        "Imports", "partner", "All other countries", top_markets)
    pt.write_bot(book)

    os.makedirs(out_dir, exist_ok=True)
    fname = "PERFORMANCE TABLES %s %d.xlsx" % (
        spec["span_short"].upper().replace("-", " "), yr)
    path = os.path.join(out_dir, fname)
    book.save(path)
    print("Saved:", path)
    return path, spec


def main():
    ap = argparse.ArgumentParser(
        description="Build quarter/semi-annual/annual performance workbooks "
                    "from raw KNBS monthly data.")
    ap.add_argument("--data-dir", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--period", choices=("quarter", "semi", "year"),
                    required=True)
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--prev-year", type=int, default=None,
                    help="comparison year (default: year - 1); the same "
                         "sub-period is compared, e.g. Apr-Jun 2026 vs "
                         "Apr-Jun 2025")
    ap.add_argument("--half", type=int, default=1, choices=(1, 2))
    ap.add_argument("--quarter", type=int, default=1)
    ap.add_argument("--top-products", type=int, default=20)
    ap.add_argument("--top-markets", type=int, default=25)
    a = ap.parse_args()
    build(a.data_dir, a.out_dir, a.period, a.year, half=a.half,
          quarter=a.quarter, top_products=a.top_products,
          top_markets=a.top_markets, prev_year=a.prev_year)


if __name__ == "__main__":
    main()
