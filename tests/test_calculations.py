"""P0 calculation tests: the numeric rules behind every KEPROBA table.

Covers, per the review's priority list:
  * USD thousand -> USD billion (Tables 1-4) and -> USD Million (Tables 5-6)
  * ranking by latest-year value
  * top-N selection
  * retaining Kenya when she falls outside the displayed rows
  * "All other ..." = total - sum(shown), None-safe
  * share = value / total with edge cases
  * trade balance = exports - imports (formula + cached value + re-read)
  * growth / CAGR
  * quarterly comparison tables (change, %, all-other, grand total)
  * review-period month selection (quarter / semi / year)
"""

import openpyxl
import pandas as pd
import pytest

import make_quarterly_tables as mqt
import make_services_tables as mst
import make_tables
from generate_report import parse_trade_balance
import period_data
from period_data import _read_one, data_sheets, load_monthly, period_spec
from pipeline_errors import PipelineError
from xlsx_tools import _finalize


# ---------------------------------------------------------------------------
# Unit conversion + ranking + top-N + Kenya retention + all-other (prepare)
# ---------------------------------------------------------------------------
def _mk(label, vals):
    return {"label": label, "vals": list(vals)}


def test_prepare_scales_thousand_to_billion_tables_1_to_4():
    # raw Trade Map downloads are USD Thousand; Tables 1-4 divide by 1e6
    ext = {
        "total": _mk("World", [2_000_000, 3_000_000]),
        "items": [_mk("Uganda", [1_250_000, 1_500_000])],
    }
    d = make_tables.prepare(ext, [2024, 2025], top_n=5, divisor=1e6,
                            is_markets=True)
    assert d["shown"][0]["vals"] == [1.25, 1.5]
    assert d["total"]["vals"] == [2.0, 3.0]


def test_prepare_scales_thousand_to_million_tables_5_6():
    # Kenya's reviewed example: exports 616.213 (USD Million) for 616,213
    ext = {
        "total": _mk("Total", [616_213, 700_000]),
        "items": [_mk("0901 Coffee", [616_213, 700_000])],
    }
    d = make_tables.prepare(ext, [2024, 2025], top_n=5, divisor=1e3,
                            is_markets=False)
    assert d["shown"][0]["vals"][0] == pytest.approx(616.213)
    assert d["total"]["vals"][1] == pytest.approx(700.0)


def test_prepare_ranks_by_latest_year_and_truncates_to_top_n():
    ext = {
        "total": _mk("World", [10_000_000] * 3),
        "items": [
            _mk("Small", [1_000_000, 1_000_000, 1_000_000]),
            _mk("Large", [9_000_000, 9_000_000, 9_000_000]),
            _mk("Mid", [5_000_000, 5_000_000, 5_000_000]),
            _mk("Tiny", [100_000, 100_000, 100_000]),
        ],
    }
    d = make_tables.prepare(ext, [2023, 2024, 2025], top_n=3,
                            divisor=1e6, is_markets=False)
    assert [it["label"] for it in d["shown"]] == ["Large", "Mid", "Small"]
    assert [it["rank"] for it in d["shown"]] == [1, 2, 3]
    # every item (even those below the cut) carries its rank
    ranks = {it["label"]: it["rank"] for it in ext["items"]}
    assert ranks["Tiny"] == 4


def test_prepare_keeps_kenya_when_outside_top_n():
    ext = {
        "total": _mk("World", [10_000_000] * 2),
        "items": [
            _mk("Uganda", [100_000, 200_000]),
            _mk("Large", [9_000_000, 9_000_000]),
            _mk("Mid", [5_000_000, 5_000_000]),
            _mk("Small", [3_000_000, 3_000_000]),
        ],
    }
    d = make_tables.prepare(ext, [2024, 2025], top_n=3, divisor=1e6,
                            is_markets=True)
    labels = [it["label"] for it in d["shown"]]
    assert labels[:3] == ["Large", "Mid", "Small"]
    assert "Kenya" not in labels           # there is no Kenya item here
    # ... but when Kenya IS present below the cut she is appended
    ext2 = {
        "total": _mk("World", [10_000_000] * 2),
        "items": [
            _mk("Kenya", [100_000, 200_000]),
            _mk("Large", [9_000_000, 9_000_000]),
            _mk("Mid", [5_000_000, 5_000_000]),
            _mk("Small", [3_000_000, 3_000_000]),
        ],
    }
    d2 = make_tables.prepare(ext2, [2024, 2025], top_n=3, divisor=1e6,
                             is_markets=True)
    labels2 = [it["label"] for it in d2["shown"]]
    assert labels2 == ["Large", "Mid", "Small", "Kenya"]
    kenya = [it for it in d2["shown"] if it["label"] == "Kenya"][0]
    assert kenya["rank"] == 4               # keeps her true rank


def test_prepare_none_latest_value_ranks_last():
    ext = {
        "total": _mk("World", [10_000_000, 10_000_000]),
        "items": [
            _mk("HasValue", [1_000_000, 1_000_000]),
            _mk("Missing", [1_000_000, None]),
        ],
    }
    d = make_tables.prepare(ext, [2024, 2025], top_n=1, divisor=1e6,
                            is_markets=False)
    assert d["shown"][0]["label"] == "HasValue"
    ranks = {it["label"]: it["rank"] for it in ext["items"]}
    assert ranks["Missing"] == 2


def test_all_other_is_total_minus_shown_per_year():
    shown = [_mk("A", [10.0, 20.0]), _mk("B", [5.0, None])]
    total = _mk("World", [100.0, 120.0])
    out = make_tables.all_other_vals(total, shown, 2)
    assert out == [85.0, 100.0]             # None treated as 0 in the sum
    # no total -> all None
    assert make_tables.all_other_vals(None, shown, 2) == [None, None]
    # total year missing -> None for that year
    assert make_tables.all_other_vals(_mk("W", [None, 120.0]), shown, 2) == \
        [None, 100.0]


def test_share_edge_cases():
    assert make_tables.share(5, 200) == pytest.approx(0.025)
    assert make_tables.share(None, 200) is None
    assert make_tables.share(5, None) is None
    assert make_tables.share(5, 0) is None
    assert make_tables.share(7, 7) == 1.0


# ---------------------------------------------------------------------------
# Trade balance (Figure 1): formula, cached value, round-trip parse
# ---------------------------------------------------------------------------
def test_trade_balance_formula_cache_and_parse(tmp_path):
    years = [2024, 2025]
    exports = [1248.907, 1300.5]
    imports = [355.622, 400.25]
    wb = openpyxl.Workbook()
    ws = wb.active
    cache = []
    make_tables.write_balance(ws, "Kenya", "USA", years, exports, imports,
                              cache=cache, skip_chart=True)
    # live Excel formula: Balance of Trade = Exports - Imports
    assert ws["B6"].value == "=B4-B5"
    assert ws["C6"].value == "=C4-C5"
    # ... with the numeric result cached for data_only readers
    assert dict(cache)["B6"] == pytest.approx(893.285)   # 1248.907 - 355.622
    assert dict(cache)["C6"] == pytest.approx(900.25)

    path = str(tmp_path / "Figure 1 Trade Balance.xlsx")
    _finalize(ws, path, cache)

    parsed = parse_trade_balance(path)
    assert parsed["years"] == years
    assert parsed["exports"] == pytest.approx(exports)
    assert parsed["imports"] == pytest.approx(imports)
    assert parsed["balance"] == pytest.approx([893.285, 900.25])


def test_balance_years_trailing_consecutive_run():
    assert make_tables._balance_years([]) == []
    assert make_tables._balance_years(list(range(2016, 2026))) == \
        list(range(2016, 2026))
    # a gap stops the run (2017 missing)
    assert make_tables._balance_years([2016, 2018, 2019, 2020]) == \
        [2018, 2019, 2020]
    # capped at n
    assert make_tables._balance_years(list(range(2000, 2026)), n=10) == \
        list(range(2016, 2026))


def test_merge_balance_series_updates_and_keeps_history():
    existing = {"years": [2020, 2021], "exports": [1.0, 2.0],
                "imports": [3.0, 4.0]}
    ys, exp, imp = make_tables.merge_balance_series(
        existing, [2021, 2022], [20.0, 9.0], [30.0, 8.0])
    assert ys == [2020, 2021, 2022]
    assert exp == [1.0, 20.0, 9.0]          # 2021 overwritten with new data
    assert imp == [3.0, 30.0, 8.0]
    # no existing series -> straight passthrough
    assert make_tables.merge_balance_series(
        None, [2025], [5.0], [6.0]) == ([2025], [5.0], [6.0])


# ---------------------------------------------------------------------------
# Growth / CAGR (services tables; used by every product-style sheet)
# ---------------------------------------------------------------------------
def test_calc_growth_rates():
    items = [{"vals": [100.0, 110.0, 121.0]}]
    mst.calc_growth_rates(items, [2020, 2021, 2022])
    assert items[0]["growth"] == pytest.approx(0.10)      # (121-110)/110
    assert items[0]["cagr"] == pytest.approx(0.10)        # (121/100)^(1/2)-1

    # fewer than two years: growth is undefined, never a fake 0
    items1 = [{"vals": [100.0]}]
    mst.calc_growth_rates(items1, [2022])
    assert items1[0]["growth"] is None
    assert items1[0].get("cagr") is None

    # missing latest value -> None growth, not zero
    items2 = [{"vals": [100.0, None]}]
    mst.calc_growth_rates(items2, [2021, 2022])
    assert items2[0]["growth"] is None


def test_calc_cagr():
    assert mst.calc_cagr([100.0, 110.0, 121.0],
                         [2020, 2021, 2022]) == pytest.approx(0.10)
    # gap years use the true year span, not the cell count
    assert mst.calc_cagr([100.0, 161.051], [2020, 2025]) == \
        pytest.approx(0.10, abs=1e-4)                      # 1.1^5 = 1.61051
    assert mst.calc_cagr([], [2020, 2021]) is None
    assert mst.calc_cagr([100.0, 90.0], [2021, 2020]) is None  # n_years <= 0
    assert mst.calc_cagr([None, None], [2020, 2021]) is None


def test_services_rank_and_top_and_all_other():
    items = [_mk("A", [1.0, 5.0]), _mk("B", [9.0, 9.0]), _mk("C", [2.0, 2.0])]
    shown = mst.rank_and_top(items, 2)
    assert [it["label"] for it in shown] == ["B", "A"]
    assert [it["rank"] for it in items] == [1, 2, 3]       # ranks on all
    assert mst.all_other_vals(_mk("T", [12.0, 16.0]), shown, 2) == [2.0, 2.0]
    assert mst.share(None, 10.0) is None
    assert mst.share(3.0, 12.0) == pytest.approx(0.25)


def test_rank_by_region_groups_and_caps():
    items = [
        {"label": "Uganda", "region": "Africa", "vals": [1.0, 5.0]},
        {"label": "Tanzania", "region": "Africa", "vals": [1.0, 9.0]},
        {"label": "Germany", "region": "Europe", "vals": [1.0, 7.0]},
        {"label": "France", "region": "Europe", "vals": [1.0, 8.0]},
        {"label": "Italy", "region": "Europe", "vals": [1.0, 6.0]},
    ]
    out = mst.rank_by_region(items, [2024, 2025], top_n=2)
    assert [it["label"] for it in out["Africa"]] == ["Tanzania", "Uganda"]
    assert [it["label"] for it in out["Europe"]] == ["France", "Germany"]
    assert out["Europe"][0]["region_rank"] == 1
    assert out["Africa"][0]["region_rank"] == 1


# ---------------------------------------------------------------------------
# Quarterly comparison tables (change, % change, all-other, grand total)
# ---------------------------------------------------------------------------
def test_rank_table_comparison_change_pct_all_other_grand():
    cur = pd.DataFrame(
        {4: [10.0, 5.0], 5: [10.0, 5.0], 6: [10.0, 5.0],
         "__total__": [30.0, 15.0]},
        index=["Uganda", "Tanzania"])
    prev = pd.DataFrame(
        {4: [8.0, 5.0], 5: [8.0, 5.0], 6: [8.0, 5.0],
         "__total__": [24.0, 15.0]},
        index=["Uganda", "Tanzania"])

    rows, all_other, grand = mqt.rank_table_comparison(
        cur, prev, months=[4, 5, 6], top_n=1, all_label="All others")

    ug = rows[0]
    assert ug["name"] == "Uganda" and ug["rank"] == 1
    assert ug["cur_total"] == pytest.approx(30.0)
    assert ug["prev_total"] == pytest.approx(24.0)
    assert ug["change"] == pytest.approx(6.0)
    assert ug["pct"] == pytest.approx(0.25)

    # all-other = grand - shown, per month and in total
    assert all_other["cur_months"] == pytest.approx([5.0, 5.0, 5.0])
    assert all_other["cur_total"] == pytest.approx(15.0)
    assert all_other["change"] == pytest.approx(0.0)

    assert grand["cur_total"] == pytest.approx(45.0)
    assert grand["prev_total"] == pytest.approx(39.0)


def test_rank_table_comparison_missing_prev_year_item():
    cur = pd.DataFrame({4: [10.0], 5: [10.0], 6: [10.0],
                        "__total__": [30.0]}, index=["Newcomer"])
    prev = pd.DataFrame({4: [0.0], 5: [0.0], 6: [0.0],
                        "__total__": [0.0]}, index=["Other"])
    rows, _ao, _g = mqt.rank_table_comparison(
        cur, prev, months=[4, 5, 6], top_n=5, all_label="All others")
    assert rows[0]["prev_total"] == 0.0
    assert rows[0]["change"] == pytest.approx(30.0)
    assert rows[0]["pct"] is None          # no previous total -> no fake 0%


def test_quarter_label():
    assert mqt.quarter_label([4, 5, 6]) == "April-June"
    assert mqt.quarter_label([1, 2, 3]) == "January-March"
    assert mqt.quarter_label([1, 4]) == "January-April"


# ---------------------------------------------------------------------------
# Review-period month selection (quarter / semi / annual)
# ---------------------------------------------------------------------------
def test_period_spec_quarter_semi_year():
    q2 = period_spec("quarter", 2026, quarter=2)
    assert q2["months"] == [4, 5, 6]
    assert q2["span"] == "April-June"
    assert q2["prev_year"] == 2025
    assert q2["span_short"] == "Apr-Jun"

    h2 = period_spec("semi", 2026, half=2)
    assert h2["months"] == list(range(7, 13))
    assert h2["span"] == "July-December"
    assert [name for name, _ in h2["subs"]] == ["Q3", "Q4"]

    yr = period_spec("year", 2026)
    assert yr["months"] == list(range(1, 13))
    assert yr["span"] == "January-December"

    with pytest.raises(PipelineError):
        period_spec("quarter", 2026, quarter=5)
    with pytest.raises(PipelineError):
        period_spec("semi", 2026, half=3)
    with pytest.raises(PipelineError):
        period_spec("fortnight", 2026)


# ---------------------------------------------------------------------------
# KNBS Flow handling: total exports = domestic exports + re-exports
# ---------------------------------------------------------------------------
def _write_knbs_csv(path, rows):
    path.write_text(
        "Year,Month,Partner,Partner.1,HS_C,Value,Flow\n"
        + "\n".join(",".join(str(c) for c in r) for r in rows) + "\n",
        encoding="utf-8")
    return str(path)


def test_knbs_flow_totals_exports_plus_reexports(tmp_path):
    csv = _write_knbs_csv(tmp_path / "knbs.csv", [
        (2026, 4, "FRANCE", "France", "0901", 2_000_000_000, "Domestic exports"),
        (2026, 4, "FRANCE", "France", "0901", 1_000_000_000, "Re-exports"),
        (2026, 5, "USA", "United States", "6203", 3_000_000_000, "Domestic exports"),
        (2026, 5, "USA", "United States", "6203", 500_000_000, "Re-exports"),
        (2026, 4, "CHINA", "China", "8471", 7_000_000_000, "Imports"),
    ])
    m = _read_one(csv, None)
    # both export flows share Side="Exports"; only imports stay separate
    sides = sorted(m["Side"].unique())
    assert sides == ["Exports", "Imports"]

    spec = period_spec("quarter", 2026, quarter=2)   # April-June
    sub = m[m["Month"].isin(spec["months"])]
    total_exports = sub[sub["Side"] == "Exports"]["Value"].sum()
    total_imports = sub[sub["Side"] == "Imports"]["Value"].sum()
    assert total_exports == pytest.approx(6.5)       # 2+1 domestic+re-exp + 3+0.5
    assert total_imports == pytest.approx(7.0)

    # the workbook DATA sheets carry the same combined totals per month
    sheets = data_sheets(sub)
    dp = sheets["DATA PARTNER"]
    apr_france = dp[(dp["Month"] == 4) & (dp["Partner"] == "FRANCE")
                    & (dp["Side"] == "Exports")]
    assert float(apr_france["Value"].iloc[0]) == pytest.approx(3.0)


def test_knbs_flow_case_insensitive_mapping(tmp_path):
    csv = _write_knbs_csv(tmp_path / "case.csv", [
        (2026, 4, "X", "X", "01", 1_000_000_000, "domestic exports"),
        (2026, 4, "X", "X", "01", 2_000_000_000, "RE-EXPORTS"),
        (2026, 4, "X", "X", "01", 4_000_000_000, "IMPORTS"),
    ])
    m = _read_one(csv, None)
    counts = m.groupby("Side")["Value"].sum().to_dict()
    assert counts["Exports"] == pytest.approx(3.0)
    assert counts["Imports"] == pytest.approx(4.0)


# ---------------------------------------------------------------------------
# KNBS Flow hardening: blank/garbage Flow rows are dropped, never guessed
# ---------------------------------------------------------------------------
def test_knbs_flow_blank_and_garbage_rows_dropped(tmp_path):
    csv = _write_knbs_csv(tmp_path / "dirty.csv", [
        (2026, 4, "FRANCE", "France", "0901", 2_000_000_000, "Domestic exports"),
        (2026, 4, "FRANCE", "France", "0901", 1_000_000_000, "Re-exports"),
        (2026, 4, "X", "X", "01", 9_000_000_000, ""),          # blank
        (2026, 4, "X", "X", "01", 8_000_000_000, "banana"),    # garbage
        (2026, 4, "X", "X", "01", 7_000_000_000, "n/a"),
        (2026, 4, "CHINA", "China", "8471", 5_000_000_000, "Imports"),
        (2026, 4, "X", "X", "01", 6_000_000_000, "???"),       # garbage
    ])
    period_data.WARNINGS.clear()
    m = _read_one(csv, None)
    # only the recognised rows survive; the junk rows are gone
    assert len(m) == 3
    exp = m[m["Side"] == "Exports"]["Value"].sum()
    imp = m[m["Side"] == "Imports"]["Value"].sum()
    assert exp == pytest.approx(3.0)      # 2 + 1 -- NOT inflated by junk
    assert imp == pytest.approx(5.0)
    # the drop is reported, with the offending values (blanks included)
    assert len(period_data.WARNINGS) == 1
    warn = period_data.WARNINGS[0]
    assert "dropped 4 row(s)" in warn
    assert "'banana'" in warn and "'???'" in warn and "'(blank)'" in warn


def test_knbs_flow_all_rows_unknown_returns_none(tmp_path):
    csv = _write_knbs_csv(tmp_path / "junk.csv", [
        (2026, 4, "X", "X", "01", 1_000_000_000, ""),
        (2026, 4, "X", "X", "01", 2_000_000_000, "mystery"),
    ])
    period_data.WARNINGS.clear()
    assert _read_one(csv, None) is None    # unusable, not exports-by-default
    assert any("dropped 2 row(s)" in w for w in period_data.WARNINGS)


def test_knbs_flow_recognised_variants_still_map(tmp_path):
    csv = _write_knbs_csv(tmp_path / "variants.csv", [
        (2026, 4, "X", "X", "01", 1_000_000_000, "Export"),
        (2026, 4, "X", "X", "01", 2_000_000_000, "Exports"),
        (2026, 4, "X", "X", "01", 4_000_000_000, "reexports"),
        (2026, 4, "X", "X", "01", 8_000_000_000, "re exports"),
        (2026, 4, "X", "X", "01", 16_000_000_000, "Domestic Export"),
        (2026, 4, "X", "X", "01", 32_000_000_000, "Import"),
    ])
    period_data.WARNINGS.clear()
    m = _read_one(csv, None)
    counts = m.groupby("Side")["Value"].sum().to_dict()
    assert counts["Exports"] == pytest.approx(31.0)
    assert counts["Imports"] == pytest.approx(32.0)
    assert period_data.WARNINGS == []      # nothing dropped


def test_load_monthly_cache_restores_flow_warnings(tmp_path):
    _write_knbs_csv(tmp_path / "knbs.csv", [
        (2026, 4, "FRANCE", "France", "0901", 2_000_000_000, "Domestic exports"),
        (2026, 4, "X", "X", "01", 9_000_000_000, "garbage-flow"),
    ])
    period_data.WARNINGS.clear()
    df1 = load_monthly(str(tmp_path))
    assert len(df1) == 1
    assert any("dropped 1 row(s)" in w for w in period_data.WARNINGS)
    # second load hits the disk cache -- warnings must survive
    period_data.WARNINGS.clear()
    df2 = load_monthly(str(tmp_path))
    assert len(df2) == 1
    assert any("dropped 1 row(s)" in w for w in period_data.WARNINGS)
