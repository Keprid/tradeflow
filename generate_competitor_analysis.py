#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_competitor_analysis.py
===============================

Competitor analysis for the Goods Trade Flow pipeline.

Reads the per-product "supplying markets" Trade Map downloads -- one workbook
per HS product, e.g. ``nigerias-imports-from-world-by-exporter_0902.xlsx``
("List of supplying markets for a product imported by Nigeria", per product)
-- and produces the KEPROBA competitor-analysis table:

    Rank | Product code | Product label | Value imported in <year>
    | Rank | Top exporters | Value exported in <year> | Share in <year> |
    Tariff Applied %

Deliverables (matching the picture template):

    COMPETITOR ANALYSIS.png   paste-ready table image (the report embeds its
                              tables as pictures)
    COMPETITOR ANALYSIS.docx  editable Word table
    COMPETITOR ANALYSIS.xlsx  editable Excel workbook

Conventions (as in the template picture):
  * products are ranked 1..n by the value imported by the partner market in
    the review year (descending);
  * each product shows its top exporters by value with their rank, value and
    share of the partner's imports of that product;
  * Kenya's row is always shown and printed in red: when she is outside the
    top N her row replaces the last slot and carries her true rank;
  * the "Tariff Applied %" column is left empty.

This module is also called from ``generate_report.py`` (Goods Trade Flow):
when such per-product files are present in ``--excel-dir`` the report gains a
"Competitor Analysis" section after Export Potential.

Usage:
    python generate_competitor_analysis.py --data TEKTAT --out TEKTAT
"""

import argparse
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

from country_names import display_name, fix_label  # noqa: E402

KENYA_RED = "FF0000"
GRID = "808080"

# Trade Map naming: "<partner>-imports-from-world-by-exporter_<HS>[ (n)].xlsx"
_SUPPLIER_RE = re.compile(
    r"imports-from-world-by-exporter[_\s-]*'?(\d{4,6})(?:\s*\(\d+\))?\.xlsx$",
    re.IGNORECASE)
_YEAR_HDR_RE = re.compile(r"^\s*'?(\d{4})\b")


def is_supplier_download(filename):
    """True when ``filename`` is a per-product supplying-markets download."""
    return bool(_SUPPLIER_RE.search(os.path.basename(str(filename))))


# ---------------------------------------------------------------------------
# Input
# ---------------------------------------------------------------------------
def find_product_supplier_files(data_dir):
    """All per-product supplying-markets workbooks under ``data_dir``.

    Returns a list of ``(path, hs_code)`` sorted by HS code, one entry per
    product code (duplicate downloads -- e.g. ``_0902 (1).xlsx`` -- are
    collapsed, the longer workbook winning).
    """
    hits = {}
    for root, _dirs, names in os.walk(data_dir):
        for name in names:
            m = _SUPPLIER_RE.search(name)
            if not m:
                continue
            path = os.path.join(root, name)
            code = m.group(1)
            prev = hits.get(code)
            if prev is None or _row_count(path) > _row_count(prev[0]):
                hits[code] = (path, code)
    return [hits[c] for c in sorted(hits)]


def _row_count(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        n = ws.max_row or 0
        wb.close()
        return n
    except Exception:
        return 0


def _header_map(ws):
    """Locate the header row of a Trade Map 'Data' sheet and map columns."""
    for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True),
                            start=1):
        labels = [str(v).strip().lower() if v is not None else ""
                  for v in row]
        if "productcd" in labels and "partnerlabel" in labels:
            cols = {lab: i for i, lab in enumerate(labels) if lab}
            years = {}
            for i, v in enumerate(row):
                if v is None:
                    continue
                m = _YEAR_HDR_RE.match(str(v))
                if m:
                    years[int(m.group(1))] = i
            return ri, cols, years
    raise ValueError("no Trade Map header row found")


def load_products(files, year=None):
    """Parse the per-product workbooks into product records.

    Each record: ``code``, ``label``, ``world`` (partner's imports of the
    product from the world), ``exporters`` (sorted by value, each with
    ``rank``/``label``/``value``/``share``/``kenya``) and ``source_file``.
    ``year`` is the review year; when omitted the latest year column is used.
    """
    products = []
    for path, code in files:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        _hdr_row, cols, years = _header_map(ws)
        if not years:
            raise ValueError("no year columns found in %s"
                             % os.path.basename(path))
        rev = year if (year in years) else max(years)
        ycol = years[rev]
        c_partner = cols.get("partnerlabel")
        c_code = cols.get("productcd")
        c_label = cols.get("productlabel")
        c_pcode = cols.get("partnercd")
        rows_by_product = {}
        for row in ws.iter_rows(min_row=_hdr_row + 1, values_only=True):
            if row is None or all(v is None for v in row):
                continue
            pcode = str(row[c_code]).strip() if row[c_code] is not None else code
            val = row[ycol]
            val = float(val) if isinstance(val, (int, float)) else None
            entry = {
                "partner": str(row[c_partner]).strip() if row[c_partner] is not None else "",
                "partner_cd": str(row[c_pcode]).strip() if row[c_pcode] is not None else "",
                "label": str(row[c_label]).strip() if row[c_label] is not None else "",
                "value": val,
            }
            rows_by_product.setdefault(pcode, []).append(entry)
        wb.close()
        for pcode, entries in rows_by_product.items():
            products.append(_product_record(pcode, entries, rev, path))
    products.sort(key=lambda p: -(p["world"] or 0.0))
    for i, p in enumerate(products, start=1):
        p["rank"] = i
    return products


def _is_world(entry):
    return entry["partner_cd"] == "000" or \
        entry["partner"].strip().lower() in ("world", "world total")


def _is_kenya(entry):
    return "kenya" in entry["partner"].strip().lower()


def _product_record(pcode, entries, year, path):
    world_rows = [e for e in entries if _is_world(e)]
    exporters = [e for e in entries
                 if not _is_world(e) and (e["value"] or 0.0) > 0.0]
    exporters.sort(key=lambda e: -(e["value"] or 0.0))
    world = world_rows[0]["value"] if world_rows else None
    if world is None:
        world = sum((e["value"] or 0.0) for e in exporters) or None
    rows = []
    for i, e in enumerate(exporters, start=1):
        share = (e["value"] / world) if world else None
        rows.append({"rank": i,
                     "label": display_name(fix_label(e["partner"])),
                     "value": e["value"], "share": share,
                     "kenya": _is_kenya(e)})
    label = next((e["label"] for e in entries if e["label"]), "")
    return {"code": pcode, "label": label, "year": year,
            "world": world, "exporters": rows, "source_file": path}


# ---------------------------------------------------------------------------
# Table model
# ---------------------------------------------------------------------------
HEADER = ["Rank in {Y}", "Product code", "Product label",
          "Value imported in {Y}", "Rank in {Y}", "Top exporters",
          "Value exported in {Y}", "Share in {Y} %", "Tariff Applied %"]


def product_blocks(products, top=5):
    """One block per product: ``top`` exporter rows, with Kenya guaranteed.

    When Kenya is outside the top ``top`` her row replaces the last slot and
    keeps her true rank.
    """
    blocks = []
    for p in products:
        rows = list(p["exporters"][:top])
        kenya = next((e for e in p["exporters"] if e["kenya"]), None)
        if kenya is not None and not any(r["kenya"] for r in rows):
            if rows:
                rows = rows[:top - 1] + [kenya]
            else:
                rows = [kenya]
        blocks.append({"product": p, "rows": rows})
    return blocks


def _fmt_value(v):
    return "" if v is None else "{:,.0f}".format(v)


def _fmt_share(s):
    return "" if s is None else "{:.1f}%".format(s * 100.0)


def _row_cells(prod, row, year):
    return [("Rank in %d" % year, row["rank"]),
            ("Product code", "'%s" % prod["code"]),
            ("Product label", _shorten(prod["label"], 95)),
            ("Value imported in %d" % year, _fmt_value(prod["world"])),
            ("Rank in %d" % year, row["rank"]),
            ("Top exporters", row["label"]),
            ("Value exported in %d" % year, _fmt_value(row["value"])),
            ("Share in %d %% " % year, _fmt_share(row["share"])),
            ("Tariff Applied %", "")]


def _shorten(text, n):
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[:n - 1].rstrip() + "\u2026"


def narrative_bullets(products, partner, top=5):
    """Short auto-generated reading notes for the report."""
    out = []
    blocks = product_blocks(products, top=top)
    lead = []
    for b in blocks:
        rows = b["rows"]
        if rows and rows[0]["kenya"]:
            lead.append(b["product"]["code"])
    if lead:
        out.append("Kenya is the leading supplier of %s to %s in the review "
                   "year (HS %s)." % (_product_phrase(lead), partner,
                                      ", ".join(lead)))
    ranked = []
    for b in blocks:
        kenya = next((r for r in b["rows"] if r["kenya"]), None)
        if kenya and kenya["share"] is not None:
            ranked.append((kenya["share"], b["product"], kenya))
    if ranked:
        ranked.sort(key=lambda t: -t[0])
        share, prod, kenya = ranked[0]
        out.append("Kenya's strongest position is HS %s (%s), where it "
                   "supplies %.1f%% of %s's imports (rank %d)."
                   % (prod["code"], _shorten(prod["label"], 45), share * 100.0,
                      partner, kenya["rank"]))
    biggest = max(blocks, key=lambda b: b["product"]["world"] or 0.0) \
        if blocks else None
    if biggest:
        p = biggest["product"]
        kenya = next((r for r in biggest["rows"] if r["kenya"]), None)
        out.append("The largest import market in the set is HS %s (%s), worth "
                   "%s USD thousand; Kenya %s." %
                   (p["code"], _shorten(p["label"], 45), _fmt_value(p["world"]),
                    ("ranks %d with %.1f%%" % (kenya["rank"],
                                               (kenya["share"] or 0.0) * 100.0))
                    if kenya else "is not among the suppliers listed"))
    return out


def _product_phrase(codes):
    return "the product" if len(codes) == 1 else "the products"


# ---------------------------------------------------------------------------
# PNG (paste-ready picture, as in the template)
# ---------------------------------------------------------------------------
_PNG_WIDTHS = [0.055, 0.065, 0.235, 0.100, 0.055, 0.150, 0.105, 0.075, 0.070]


def render_png(products, out_png, year=None, partner_label="", top=5):
    """Render the competitor-analysis table as a PNG image."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    import textwrap

    year = year or (products[0]["year"] if products else "")
    blocks = product_blocks(products, top=top)
    headers = [h.format(Y=year) for h in HEADER]

    def wrap(text, chars):
        lines = textwrap.wrap(str(text), chars, break_long_words=False,
                              break_on_hyphens=False) or [""]
        return lines

    widths = _PNG_WIDTHS
    total = sum(widths)
    widths = [w / total for w in widths]

    row_h = 0.24
    wrap_chars = [3, 5, 34, 9, 3, 16, 9, 6, 6]
    # per-block height driven by the tallest wrapped label
    block_metrics = []
    for b in blocks:
        label_lines = len(wrap(_shorten(b["product"]["label"], 95),
                               wrap_chars[2]))
        name_lines = max([len(wrap(r["label"], wrap_chars[5]))
                          for r in b["rows"]] or [1])
        lines = max(label_lines, name_lines)
        block_h = max(row_h * len(b["rows"]),
                      0.20 * lines + 0.06 * (len(b["rows"]) - 1))
        block_metrics.append(block_h)

    fig_h = 0.34 + sum(block_metrics) + 0.18
    fig_w = 10.2
    dpi = 220
    fig = plt.figure(figsize=(fig_w, fig_h), dpi=dpi)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, fig_w)
    ax.set_ylim(0, fig_h)
    ax.axis("off")

    font = _serif()
    xs = [0.0]
    for w in widths:
        xs.append(xs[-1] + w)

    def cell_rect(c, y_top, h):
        x0 = xs[c] * fig_w
        w = (xs[c + 1] - xs[c]) * fig_w
        y0 = fig_h - y_top - h
        return x0, y0, w, h

    def draw_cell(c, y_top, h, text, bold=False, red=False,
                  align="left", valign="center", size=6.6):
        x0, y0, w, h = cell_rect(c, y_top, h)
        ax.add_patch(Rectangle((x0, y0), w, h, fill=False,
                               edgecolor="black", linewidth=0.5))
        if text == "" or text is None:
            return
        lines = wrap(text, wrap_chars[c])
        n = len(lines)
        lh = size / 72.0 * 1.25
        block_h = lh * n
        if valign == "top":
            ty = y0 + h - 0.045 - lh * 0.85
        else:
            ty = y0 + (h + block_h) / 2.0 - lh * 0.85
        tx = x0 + 0.035
        ha = "left"
        if align == "center":
            tx = x0 + w / 2.0
            ha = "center"
        elif align == "right":
            tx = x0 + w - 0.035
            ha = "right"
        color = "#%s" % KENYA_RED if red else "black"
        for line in lines:
            ax.text(tx, ty, line, ha=ha, va="baseline", fontsize=size,
                    fontfamily=font, fontweight="bold" if bold else "normal",
                    color=color)
            ty -= lh

    aligns = ["center", "center", "left", "right", "center", "left",
              "right", "right", "right"]

    # header
    y = 0.0
    h0 = 0.34
    for c, text in enumerate(headers):
        draw_cell(c, y, h0, text, bold=True, align="center", size=6.4)
    y += h0

    for bi, b in enumerate(blocks):
        prod = b["product"]
        rows = b["rows"]
        top_y = y
        block_h = block_metrics[bi]
        row_h_each = block_h / max(len(rows), 1)

        # product columns merged over the block
        merged_text = {0: str(prod["rank"]), 1: "'%s" % prod["code"],
                       2: _shorten(prod["label"], 95),
                       3: _fmt_value(prod["world"])}
        for c in (0, 1, 2, 3):
            x0, y0, w, hh = cell_rect(c, top_y, block_h)
            ax.add_patch(Rectangle((x0, y0), w, hh, fill=False,
                                   edgecolor="black", linewidth=0.5))
            draw_cell_merged(c, top_y, block_h, merged_text[c], aligns[c],
                             font, ax, wrap_chars[c], fig_w, xs, fig_h,
                             size=6.6,
                             valign="top" if c == 2 else "center")

        yy = top_y
        for r in rows:
            red = r["kenya"]
            vals = [r["rank"], r["label"], _fmt_value(r["value"]),
                    _fmt_share(r["share"]), ""]
            for j, c in enumerate((4, 5, 6, 7, 8)):
                draw_cell(c, yy, row_h_each, vals[j], red=red,
                          align=aligns[c], size=6.6)
            yy += row_h_each
        y += block_h

    os.makedirs(os.path.dirname(os.path.abspath(out_png)) or ".", exist_ok=True)
    fig.savefig(out_png, dpi=dpi, facecolor="white")
    plt.close(fig)
    return out_png


def draw_cell_merged(c, y_top, h, text, align, font, ax, wrap_chars, fig_w,
                     xs, fig_h, bold=False, size=6.6, valign="center"):
    """Text inside an already-bounded merged cell."""
    import textwrap
    w = (xs[c + 1] - xs[c]) * fig_w
    x0 = xs[c] * fig_w
    y0 = fig_h - y_top - h
    if text == "" or text is None:
        return
    lines = textwrap.wrap(str(text), wrap_chars, break_long_words=False,
                          break_on_hyphens=False) or [""]
    lh = size / 72.0 * 1.25
    if valign == "top":
        ty = y0 + h - 0.045 - lh * 0.85
    else:
        ty = y0 + (h + lh * len(lines)) / 2.0 - lh * 0.85
    tx = x0 + 0.035
    ha = "left"
    if align == "center":
        tx = x0 + w / 2.0
        ha = "center"
    elif align == "right":
        tx = x0 + w - 0.035
        ha = "right"
    for line in lines:
        ax.text(tx, ty, line, ha=ha, va="baseline", fontsize=size,
                fontfamily=font, fontweight="bold" if bold else "normal",
                color="black")
        ty -= lh


def _serif():
    import matplotlib.font_manager as fm
    installed = {f.name for f in fm.fontManager.ttflist}
    for name in ("Liberation Serif", "Times New Roman", "DejaVu Serif"):
        if name in installed:
            return name
    return "DejaVu Serif"


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def write_docx(products, out_docx, year=None, partner_label="", top=5):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    year = year or (products[0]["year"] if products else "")
    blocks = product_blocks(products, top=top)
    headers = [h.format(Y=year) for h in HEADER]

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.2)
    sec.top_margin = sec.bottom_margin = Cm(1.2)

    title = doc.add_paragraph("Competitor Analysis")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.runs[0] if title.runs else title.add_run()
    run.bold = True
    run.font.size = Pt(12)

    cap = doc.add_paragraph("Table: Competitor Analysis"
                            + (" - %s, %s" % (partner_label, year)
                               if partner_label else ""))
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)

    n_data_rows = sum(len(b["rows"]) for b in blocks)
    table = doc.add_table(rows=1 + n_data_rows, cols=len(HEADER))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    def set_cell(cell, text, bold=False, red=False, align="left", size=7):
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = para.add_run(str(text))
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        run.bold = bold
        if red:
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    def vmerge(cell_a, cell_b):
        tc_a = cell_a._tc
        tc_b = cell_b._tc
        tc_prop = tc_a.get_or_add_tcPr()
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), "restart")
        tc_prop.append(merge)
        tc_prop_b = tc_b.get_or_add_tcPr()
        merge_b = OxmlElement("w:vMerge")
        tc_prop_b.append(merge_b)

    aligns = ["center", "center", "left", "right", "center", "left",
              "right", "right", "right"]
    for c, text in enumerate(headers):
        set_cell(table.rows[0].cells[c], text, bold=True, align="center",
                 size=6.5)

    ri = 1
    for b in blocks:
        prod = b["product"]
        rows = b["rows"]
        start = ri
        for r in rows:
            red = r["kenya"]
            set_cell(table.rows[ri].cells[4], r["rank"], red=red,
                     align=aligns[4])
            set_cell(table.rows[ri].cells[5], r["label"], red=red,
                     align=aligns[5])
            set_cell(table.rows[ri].cells[6], _fmt_value(r["value"]), red=red,
                     align=aligns[6])
            set_cell(table.rows[ri].cells[7], _fmt_share(r["share"]), red=red,
                     align=aligns[7])
            set_cell(table.rows[ri].cells[8], "", red=red, align=aligns[8])
            ri += 1
        for j, c in enumerate((0, 1, 2, 3)):
            text = [str(prod["rank"]), "'%s" % prod["code"],
                    _shorten(prod["label"], 200),
                    _fmt_value(prod["world"])][j]
            cell = table.rows[start].cells[c]
            set_cell(cell, text, align=aligns[c])
            for rr in range(start + 1, start + len(rows)):
                vmerge(cell, table.rows[rr].cells[c])

    widths_cm = [1.15, 1.35, 5.6, 1.9, 1.15, 3.6, 2.05, 1.45, 1.5]
    for row in table.rows:
        for c, w in enumerate(widths_cm):
            row.cells[c].width = Cm(w)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Note: values are in USD thousand. Kenya's row is highlighted in "
        "red; where Kenya is outside the top exporters her row shows her "
        "true rank. Tariff data not provided in the source downloads.")
    note_run.italic = True
    note_run.font.size = Pt(8)

    src = doc.add_paragraph()
    src_run = src.add_run("Source: International Trade Centre database; "
                          "Compiled by KEPROBA")
    src_run.font.size = Pt(8)

    os.makedirs(os.path.dirname(os.path.abspath(out_docx)) or ".", exist_ok=True)
    doc.save(out_docx)
    return out_docx


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
def write_xlsx(products, out_xlsx, year=None, partner_label="", top=5):
    year = year or (products[0]["year"] if products else "")
    blocks = product_blocks(products, top=top)
    headers = [h.format(Y=year) for h in HEADER]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Competitor Analysis"

    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font = Font(name="Century Gothic", bold=True)
    body_font = Font(name="Century Gothic")
    red_font = Font(name="Century Gothic", color=KENYA_RED)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    center = Alignment(horizontal="center", vertical="center")
    aligns = [center, center, left, right, center, left, right, right, right]

    for c, text in enumerate(headers, start=1):
        cell = ws.cell(1, c, text)
        cell.font = hdr_font
        cell.border = border
        cell.alignment = center

    ri = 2
    for b in blocks:
        prod = b["product"]
        rows = b["rows"]
        start = ri
        for r in rows:
            vals = [r["rank"], "'%s" % prod["code"], None, None,
                    r["rank"], r["label"], r["value"], r["share"], None]
            # product columns filled after merge
            font = red_font if r["kenya"] else body_font
            for c, val in ((5, vals[4]), (6, vals[5]), (7, vals[6]),
                           (8, vals[7]), (9, vals[8])):
                cell = ws.cell(ri, c, val)
                cell.font = font
                cell.border = border
                cell.alignment = aligns[c - 1]
                if c in (7,) and isinstance(val, (int, float)):
                    cell.number_format = "#,##0"
                if c in (8,) and isinstance(val, (int, float)):
                    cell.number_format = "0.0%"
            ri += 1
        merged_vals = [prod["rank"], "'%s" % prod["code"], prod["label"],
                       prod["world"]]
        for j, c in enumerate((1, 2, 3, 4)):
            cell = ws.cell(start, c, merged_vals[j])
            cell.font = body_font
            cell.border = border
            cell.alignment = aligns[c - 1]
            if c == 4 and isinstance(merged_vals[j], (int, float)):
                cell.number_format = "#,##0"
            if c == 2:
                cell.number_format = "@"
            if len(rows) > 1:
                ws.merge_cells(start_row=start, start_column=c,
                               end_row=start + len(rows) - 1, end_column=c)
                for rr in range(start + 1, start + len(rows)):
                    ws.cell(rr, c).border = border

    for c, w in zip(range(1, 10), (8, 11, 46, 14, 8, 24, 15, 11, 11)):
        ws.column_dimensions[get_column_letter(c)].width = w

    nr = ws.max_row + 2
    ws.cell(nr, 1, "Note: values in USD thousand; Kenya's row in red; "
                   "where Kenya is outside the top exporters her row carries "
                   "her true rank. Tariff data not provided in the source "
                   "downloads.").font = Font(name="Century Gothic", italic=True,
                                             size=9)
    ws.cell(nr + 1, 1, "Source: International Trade Centre database; "
                       "Compiled by KEPROBA").font = Font(
        name="Century Gothic", italic=True, size=9)

    os.makedirs(os.path.dirname(os.path.abspath(out_xlsx)) or ".", exist_ok=True)
    wb.save(out_xlsx)
    return out_xlsx


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------
def generate(data_dir, out_dir, year=None, top=5, basename="COMPETITOR ANALYSIS"):
    files = find_product_supplier_files(data_dir)
    if not files:
        raise SystemExit("no per-product supplying-markets files "
                         "(*-imports-from-world-by-exporter_<HS>.xlsx) "
                         "found under %s" % data_dir)
    products = load_products(files, year=year)
    rev = products[0]["year"]
    partner = ""
    try:
        wb = openpyxl.load_workbook(files[0][0], read_only=True, data_only=True)
        ws = wb.worksheets[0]
        hdr, cols, _y = _header_map(ws)
        for row in ws.iter_rows(min_row=hdr + 1, max_row=hdr + 1, values_only=True):
            partner = str(row[cols["reporterlabel"]]).strip()
        wb.close()
    except Exception:
        pass

    os.makedirs(out_dir, exist_ok=True)
    png = render_png(products, os.path.join(out_dir, basename + ".png"),
                     year=rev, partner_label=partner, top=top)
    docx_p = write_docx(products, os.path.join(out_dir, basename + ".docx"),
                        year=rev, partner_label=partner, top=top)
    xlsx_p = write_xlsx(products, os.path.join(out_dir, basename + ".xlsx"),
                        year=rev, partner_label=partner, top=top)
    return {"products": products, "year": rev, "partner": partner,
            "png": png, "docx": docx_p, "xlsx": xlsx_p}


def main():
    ap = argparse.ArgumentParser(
        description="Generate the competitor analysis table from per-product "
                    "supplying-markets Trade Map downloads.")
    ap.add_argument("--data", required=True,
                    help="folder holding *-imports-from-world-by-exporter_<HS>.xlsx")
    ap.add_argument("--out", default=None,
                    help="output folder (default: same as --data)")
    ap.add_argument("--year", type=int, default=None,
                    help="review year (default: latest year in the files)")
    ap.add_argument("--top", type=int, default=5,
                    help="exporters shown per product (default 5)")
    ap.add_argument("--basename", default="COMPETITOR ANALYSIS",
                    help="output file basename")
    args = ap.parse_args()
    result = generate(args.data, args.out or args.data, year=args.year,
                      top=args.top, basename=args.basename)
    print("Competitor analysis (%d products, %d):" %
          (len(result["products"]), result["year"]))
    for p in result["products"]:
        print("  %d. HS %s  imported %s  (%s)" %
              (p["rank"], p["code"], _fmt_value(p["world"]),
               _shorten(p["label"], 50)))
    for key in ("png", "docx", "xlsx"):
        print("  saved: %s" % result[key])


if __name__ == "__main__":
    main()
