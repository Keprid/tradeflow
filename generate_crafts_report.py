#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_crafts_report.py
=========================

Data-driven "product profile" report generator for Kenya's commercial crafts,
built from the ITC / Trade Map *product-group* ("list of ...") downloads that
the webapp collects in one folder.

The download set is **export-focused**: one folder per run holds, for each of
the six craft categories, up to four Trade Map files (all values in USD
thousand, period 2021-2025):

  * ``List_of_exported_products_for_the_selected_product_group_(<Category>)``
        world exports by 6-digit product of the category
  * ``List_of_exporters_for_the_selected_product_group_(<Category>)``
        world exports by economy of the category
  * ``List_of_importing_markets_for_a_product_group_exported_by_Kenya``
        Kenya's exports by destination market of the category
  * ``List_of_products_exported_by_Kenya``
        Kenya's exports by 6-digit product of the category

The analysis is presented **category by category first** (Kenya's product
mix, Kenya's destination markets and the world exporters of each category)
and then for **commercial crafts as a whole** (Kenya's exports by category,
Kenya's aggregate destinations, the world exporters and the world product
trend), following the layout of the "Product Profile for Textile and apparel
2024" sample:

  * a styled Word report (.docx) with the same front matter, own tables and
    narrative bullets (no screenshots), and
  * an editable Excel deliverable ("... TABLES.xlsx") with the same tables.

Missing files are tolerated: any component that cannot be loaded simply
skips its table and logs a warning - this report is export-performance
focused and never requires import-side files.

Usage:
    python generate_crafts_report.py --data-dir crafts [--output out.docx]
                                      [--tmp output/.tmp]
"""

import argparse
import os
import re

import openpyxl

import xlsx_compat

from generate_product_profile import (
    ProfileBuilder, make_donut,
    cagr, yoy_change, display, usd_phrase, growth_phrase, yoy_phrase,
    period_phrase, ordinal_list, _shares, _year_totals, top_rows,
    _ranked_rows, short_label, _ordinal,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# role -> filename keyword used to classify an uploaded/downloaded file
ROLE_KEYWORDS = (
    ("exported_products",
     "list_of_exported_products_for_the_selected_product_group"),
    ("exporters", "list_of_exporters_for_the_selected_product_group"),
    ("destinations",
     "list_of_importing_markets_for_a_product_group_exported_by_kenya"),
    ("kenya_products", "list_of_products_exported_by_kenya"),
)

SUPPORTED_EXT = (".xlsx", ".xlsm", ".xls", ".xlsb", ".csv")


# --------------------------------------------------------------------------
# Loading
# --------------------------------------------------------------------------
def _read_rows(path):
    """Return ``(rows, product_group)`` for a spreadsheet or HTML ``.xls``.

    Rows are lists of trimmed strings.  ``product_group`` is the ITC
    "Product group:" label recovered either from the HTML metadata or from
    the tag the webapp embeds in column 20 when converting HTML downloads to
    ``.xlsx``.
    """
    rows = []
    product_group = ""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = [["" if c is None else str(c).strip() for c in r]
                for r in ws.iter_rows(values_only=True) if any(r)]
        wb.close()
        if rows and len(rows[0]) >= 20:
            tag = str(rows[0][19] or "").strip()
            if tag.startswith("Product group"):
                product_group = tag.split(":", 1)[1].strip()
    except Exception:
        rows = []
    if not rows:
        text = xlsx_compat._read_text_any(path)
        m = re.search(r"Product group\s*:\s*([^<]+)", text)
        if m:
            product_group = m.group(1).strip()
        parser = xlsx_compat.HTMLTableParser()
        parser.feed(text)
        if parser.tables:
            rows = [[str(c).strip() for c in r]
                    for r in max(parser.tables, key=len)]
    return rows, product_group


def _year_columns(header):
    """(years, column indexes) of the ``20XX`` columns in ``header``."""
    years, cols = [], []
    for i, h in enumerate(header):
        m = re.search(r"(20[12]\d)", h or "")
        if m:
            y = int(m.group(1))
            if 2000 <= y <= 2100 and y not in years:
                years.append(y)
                cols.append(i)
    return years, cols


def _num(text):
    """Parse a numeric cell; ``None`` when empty / not a number."""
    t = str(text or "").replace(",", "").replace("\xa0", " ").strip()
    if not t:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def _row_values(row, years, cols):
    return {years[i]: _num(row[c]) for i, c in enumerate(cols)}


def _clean_label(label):
    return re.sub(r"\s+", " ", str(label or "").strip())


def _clean_category(name):
    """Normalise a craft category title (filename or ITC label)."""
    name = str(name or "").replace("_", " ").strip()
    name = re.sub(r"\s+", " ", name)
    if name.lower().endswith(" new"):
        name = name[:-4].strip()
    return name


def _category_from_filename(fname):
    """Category stored in the ``(<Category>)`` suffix of the filename."""
    m = re.search(r"\(([^()]+)\)", fname)
    return _clean_category(m.group(1)) if m else ""


def _parse_grid(rows, role):
    """Split a parsed grid into ``(years, records, tot)``.

    ``records`` is a list of ``{"label", "code", "years"}`` (code only for
    product grids).  ``tot`` is the grid's own total row (the ``"World"``
    row of a country grid), if any.
    """
    years, cols = [], []
    for r in rows:
        years, cols = _year_columns(r)
        if years:
            break
    if not years:
        return [], [], None
    header_idx = rows.index(r)
    records, tot = [], None
    for row in rows[header_idx + 1:]:
        if not row or not any(row):
            continue
        values = _row_values(row, years, cols)
        if role in ("destinations", "exporters"):
            label = _clean_label(row[0])
            if not label:
                continue
            if label.lower() in ("world", "world total"):
                tot = {"label": label, "years": values}
                continue
            records.append({"label": label, "years": values})
        else:
            code = str(row[0] or "").lstrip("'").strip()
            label = _clean_label(row[1])
            low = label.lower()
            if low.startswith("total") or "all products" in low:
                tot = {"code": code, "label": label, "years": values}
                continue
            if not code and not label:
                continue
            records.append({"code": code, "label": label, "years": values})
    return years, records, tot


class CraftCategory:
    """One craft product group with its four optional data grids."""

    def __init__(self, title):
        self.title = title
        self.kenya_products = []     # [{"code","label","years"}]
        self.destinations = []       # [{"label","years"}]
        self.exporters = []          # [{"label","years"}]
        self.world_products = []     # [{"code","label","years"}]
        self.years = []
        self.kenya_product_total = None  # TOTAL row of the Kenya product grid
        self.dest_total = None       # "World" row of the Kenya destinations grid
        self.world_exp_total = None  # "World" row of the world exporters grid
        self.world_prod_total = None  # TOTAL row of the world product grid
        self.kenya_merchandise = {}  # Kenya total exports (TOTAL All products)

    @staticmethod
    def _val(row, year):
        return (row or {}).get("years", {}).get(year)

    def total_kenya(self, year):
        """Kenya's exports of the category in ``year``.

        Prefers the destination-side total (the ITC "World" row of the
        importing-markets grid) because that grid is present for every
        category; falls back to the sum of destination rows, the total row
        of the Kenya product grid, then the sum of its product rows.
        """
        v = self._val(self.dest_total, year)
        if v:
            return v
        v = self._series_total(self.destinations, year)
        if v:
            return v
        v = self._val(self.kenya_product_total, year)
        if v:
            return v
        return self._series_total(self.kenya_products, year)

    def total_dest(self, year):
        return self.total_kenya(year)

    def total_world(self, year):
        """World exports of the category's products in ``year``."""
        v = self._val(self.world_exp_total, year)
        if v:
            return v
        return self._series_total(self.exporters, year)

    @staticmethod
    def _series_total(rows, year):
        return sum((r["years"].get(year) or 0.0) for r in rows)


class CraftsData:
    """Loads every craft grid in a folder and derives the analysis tables."""

    def __init__(self, data_dir):
        self.data_dir = os.path.abspath(data_dir)
        self.warnings = []
        self.paths = {}              # (role, cat_key) -> path
        self.categories = {}         # normalized key -> CraftCategory
        self._scan()
        for cat in self.categories.values():
            cat.years = self._category_years(cat)
        self.years = self._all_years()
        self.order = self._category_order()
        if not self.categories:
            self.warnings.append(
                "no craft product-group files could be loaded from %r"
                % self.data_dir)

    # -- file scan ----------------------------------------------------------
    def _scan(self):
        for fname in sorted(os.listdir(self.data_dir)):
            path = os.path.join(self.data_dir, fname)
            if not os.path.isfile(path):
                continue
            low = fname.lower()
            if not low.endswith(SUPPORTED_EXT):
                continue
            role = self._classify(fname)
            if role is None:
                continue
            rows, pg = _read_rows(path)
            cat_key = self._category_key(fname, role, pg)
            if not cat_key:
                self.warnings.append(
                    "could not identify the product group of %s - skipped"
                    % fname)
                continue
            key = (role, cat_key)
            if key in self.paths:
                continue  # duplicate download - keep the first
            self.paths[key] = path
            years, records, tot = _parse_grid(rows, role)
            if not years:
                self.warnings.append(
                    "no yearly table found in %s - skipped" % fname)
                continue
            self._attach(cat_key, role, years, records, tot)
        if not self.categories:
            self.warnings.append(
                "no craft product-group files could be loaded from %r"
                % self.data_dir)

    @staticmethod
    def _classify(fname):
        low = fname.lower()
        for role, kw in ROLE_KEYWORDS:
            if kw in low:
                return role
        return None

    def _category_key(self, fname, role, pg):
        name = pg or _category_from_filename(fname)
        return _clean_category(name) if name else ""

    def _attach(self, cat_key, role, years, records, tot):
        cat = self.categories.setdefault(cat_key, CraftCategory(cat_key))
        # The group-name row of a product grid (no code, label == category)
        # is that grid's total line - keep it as the total, not a product.
        group_total = [r for r in records
                       if not r.get("code")
                       and _clean_label(r.get("label")) == cat_key]
        if group_total:
            records = [r for r in records if r not in group_total]
        else:
            group_total = None

        if role == "kenya_products":
            if tot and ("all products" in str(tot.get("label", "")).lower()
                        or str(tot.get("label", "")).lower().startswith(
                            "total")):
                cat.kenya_merchandise = dict(tot["years"])
            cat.kenya_product_total = group_total[0] if group_total else None
            cat.kenya_products = [r for r in records if r.get("code")]
        elif role == "destinations":
            cat.dest_total = tot or (group_total[0] if group_total else None)
            cat.destinations = records
        elif role == "exporters":
            cat.world_exp_total = tot or (group_total[0] if group_total else None)
            cat.exporters = records
        else:  # exported_products -> world exports by product
            cat.world_prod_total = group_total[0] if group_total else None
            cat.world_products = [r for r in records if r.get("code")]

    # -- derived ------------------------------------------------------------
    def _category_years(self, cat):
        year_sets = []
        for grid in (cat.kenya_products, cat.destinations, cat.exporters,
                     cat.world_products):
            for r in grid:
                year_sets.append(set(r["years"]))
        return sorted({y for s in year_sets for y in s})

    def _all_years(self):
        return sorted({y for cat in self.categories.values()
                       for y in cat.years})

    @property
    def review_year(self):
        return self.years[-1] if self.years else None

    @property
    def start_year(self):
        return self.years[0] if self.years else None

    def _category_order(self):
        rev = self.review_year
        order = sorted(self.categories.values(),
                       key=lambda c: (c.total_kenya(rev) or 0.0) if rev
                       else 0.0, reverse=True)
        return order

    def kenya_by_category(self):
        """Rows = one per category: Kenya's exports of that category."""
        rev = self.review_year
        rows = [{"label": c.title,
                 "years": {y: c.total_kenya(y) for y in self.years}}
                for c in self.order]
        return sorted(rows, key=lambda r: r["years"].get(rev) or 0.0,
                      reverse=True)

    def _aggregate(self, attr):
        """Sum one country-level grid across categories (keyed by label)."""
        out = {}
        for cat in self.categories.values():
            for r in getattr(cat, attr, []):
                lab = r["label"]
                d = out.setdefault(lab, {y: 0.0 for y in self.years})
                for y, v in r["years"].items():
                    if v:
                        d[y] = d[y] + v
        return [{"label": lab, "years": d} for lab, d in out.items()]

    def destinations(self):
        """Kenya's export destinations across all categories."""
        rows = self._aggregate("destinations")
        rev = self.review_year
        return sorted(rows, key=lambda r: r["years"].get(rev) or 0.0,
                      reverse=True)

    def exporters(self):
        """World exporters of crafts across all categories."""
        rows = self._aggregate("exporters")
        rev = self.review_year
        return sorted(rows, key=lambda r: r["years"].get(rev) or 0.0,
                      reverse=True)

    def kenya_products(self):
        """Kenya's exports of crafts by 6-digit product (aggregate)."""
        out = {}
        for cat in self.categories.values():
            for r in cat.kenya_products:
                code = r["code"]
                d = out.setdefault(code, {"code": code, "label": r["label"],
                                          "years": {y: 0.0
                                                    for y in self.years}})
                for y, v in r["years"].items():
                    if v:
                        d["years"][y] = d["years"][y] + v
        rows = list(out.values())
        rev = self.review_year
        return sorted(rows, key=lambda r: r["years"].get(rev) or 0.0,
                      reverse=True)

    def world_products(self):
        """World exports of crafts by 6-digit product (across categories)."""
        out = {}
        for cat in self.categories.values():
            for r in cat.world_products:
                code = r["code"]
                if not code:
                    continue
                d = out.setdefault(code, {"code": code, "label": r["label"],
                                          "years": {y: 0.0
                                                    for y in self.years}})
                for y, v in r["years"].items():
                    if v:
                        d["years"][y] = d["years"][y] + v
        rows = list(out.values())
        rev = self.review_year
        return sorted(rows, key=lambda r: r["years"].get(rev) or 0.0,
                      reverse=True)

    def kenya_merchandise_total(self):
        """Kenya's total merchandise exports (for specialization share).

        The TOTAL All products row repeats in several files (mirror/summary
        duplicates), so take the max value per year instead of summing.
        """
        out = {}
        for cat in self.categories.values():
            for y, v in cat.kenya_merchandise.items():
                if v and v > (out.get(y) or 0.0):
                    out[y] = v
        return out


# --------------------------------------------------------------------------
# Narrative helpers
# --------------------------------------------------------------------------
def _growth_sentence(rows, years, family):
    totals = _year_totals(rows)
    rev = years[-1]
    g = growth_phrase(cagr([totals.get(y) for y in years], years),
                      period_phrase(years[0], rev))
    yo = yoy_phrase(yoy_change([totals.get(y) for y in years], years),
                    years[-2], years[-1])
    if not g and not yo:
        return ""
    txt = "%s %s." % (family, g) if g else ""
    if yo:
        txt = (txt[:-1] + ", and %s." % yo) if txt else \
            ("Between %d and %d they %s." % (years[-2], rev,
                                             yo[len("declined by"):]
                                             if yo.startswith("declined")
                                             else yo))
    return txt


def _narrative_product(b, rows, family, years, rev, total):
    """Bullets under a Kenya-by-product table."""
    if not rows:
        return
    total = total if total is not None else \
        _year_totals(rows).get(rev)
    if total:
        b.add_bullet("Kenya's total exports of %s were %s in %d."
                     % (family, usd_phrase(total), rev))
    lead = rows[0]
    if lead and (lead["years"].get(rev) or 0.0) > 0:
        share = (lead["years"].get(rev) or 0.0) / total * 100 if total else 0
        b.add_bullet("The leading export product was %s (%s; %.1f%% of "
                     "Kenya's exports of the category)."
                     % (short_label(lead["label"], 70), lead["code"], share))
    follows = [r for r in rows[1:3] if (r["years"].get(rev) or 0.0) > 0]
    if total and follows:
        b.add_bullet("It was followed by %s."
                     % ordinal_list(
                         ["%s (%s; %.1f%%)"
                          % (short_label(r["label"], 55), r["code"],
                             (r["years"].get(rev) or 0.0) / total * 100)
                          for r in follows]))
    txt = _growth_sentence(rows, years, "Kenya's exports of %s" % family)
    if txt:
        b.add_bullet(txt)


def _narrative_dest(b, rows, family, years, rev, total):
    """Bullets under a Kenya-destination table."""
    if not rows:
        return
    total = total if total is not None else \
        _year_totals(rows).get(rev)
    if total:
        b.add_bullet("Kenya's exports of %s were %s in %d."
                     % (family, usd_phrase(total), rev))
    lead = rows[0]
    if lead and (lead["years"].get(rev) or 0.0) > 0:
        b.add_bullet("The leading destination market was %s (%s; %.1f%% "
                     "of Kenya's exports of the category)."
                     % (lead["label"], usd_phrase(lead["years"].get(rev)),
                        (lead["years"].get(rev) or 0.0) / total * 100
                        if total else 0))
    names = ["%s (%s; %.1f%%)"
             % (r["label"], usd_phrase(r["years"].get(rev)),
                (r["years"].get(rev) or 0.0) / total * 100 if total else 0)
             for r in rows[1:4] if (r["years"].get(rev) or 0.0) > 0]
    if names:
        b.add_bullet("Other leading destinations were %s."
                     % ordinal_list(names))
    txt = _growth_sentence(rows, years, "Kenya's exports of %s" % family)
    if txt:
        b.add_bullet(txt)


def _narrative_exporters(b, rows, family, years, rev, total):
    """Bullets under a world-exporter table."""
    if not rows:
        return
    total = total if total is not None else \
        _year_totals(rows).get(rev)
    lead = rows[0]
    if lead and (lead["years"].get(rev) or 0.0) > 0:
        b.add_bullet("%s was the world's leading exporter of %s in %d, "
                     "with %s (%.1f%% of the world total)."
                     % (lead["label"], family, rev,
                        usd_phrase(lead["years"].get(rev)),
                        (lead["years"].get(rev) or 0.0) / total * 100
                        if total else 0))
    names = ["%s (%s; %.1f%%)"
             % (r["label"], usd_phrase(r["years"].get(rev)),
                (r["years"].get(rev) or 0.0) / total * 100 if total else 0)
             for r in rows[1:5] if (r["years"].get(rev) or 0.0) > 0]
    if names:
        b.add_bullet("The top five exporters were %s." % ordinal_list(names))
    if total:
        top5 = sum((r["years"].get(rev) or 0.0) for r in rows[:5]) / total * 100
        b.add_bullet("Together they accounted for %.1f%% of the world total."
                     % top5)
    k_row = next((r for r in rows if r["label"] == "Kenya"), None)
    if k_row and (k_row["years"].get(rev) or 0.0) > 0:
        ranked = sorted(rows, key=lambda r: r["years"].get(rev) or 0.0,
                        reverse=True)
        rank = next((i for i, r in enumerate(ranked, 1)
                     if r["label"] == "Kenya"), None)
        if rank:
            b.add_bullet("Kenya ranked %s among the world's exporters of %s "
                         "in %d, with exports valued at %s."
                         % (_ordinal(rank), family, rev,
                            usd_phrase(k_row["years"].get(rev))))


def _narrative_world_products(b, rows, years, rev, total):
    """Bullets under a world-exports-by-product table."""
    if not rows:
        return
    total = total if total is not None else \
        _year_totals(rows).get(rev)
    lead = rows[0]
    if lead and (lead["years"].get(rev) or 0.0) > 0:
        b.add_bullet("The leading product exported globally was %s (%s; "
                     "%.1f%% of the world total)."
                     % (short_label(lead["label"], 70), lead["code"],
                        (lead["years"].get(rev) or 0.0) / total * 100
                        if total else 0))
    follows = [r for r in rows[1:3] if (r["years"].get(rev) or 0.0) > 0]
    if follows:
        b.add_bullet("It was followed by %s."
                     % ordinal_list(
                         ["%s (%s; %.1f%%)"
                          % (short_label(r["label"], 55), r["code"],
                             (r["years"].get(rev) or 0.0) / total * 100
                             if total else 0)
                          for r in follows]))
    txt = _growth_sentence(rows, years, "World exports of this category")
    if txt:
        b.add_bullet(txt)


# --------------------------------------------------------------------------
# Word builder
# --------------------------------------------------------------------------
def section_categories(b, cfg, data, source, tmp_dir):
    """Category-level analysis - one block per craft product group,
    ranked by Kenya's export value in the review year."""
    b.add_heading("KENYA'S CRAFT EXPORTS - CATEGORY ANALYSIS", level=1)
    b.add_para(
        "The export performance of commercial crafts is analysed first at "
        "the level of each of the product groups that make up the craft "
        "portfolio. For every category the profile presents Kenya's export "
        "product mix, Kenya's destination markets and the world exporters "
        "of the category's products.")
    for cat in data.order:
        rev = data.review_year
        years = data.years
        if not (cat.kenya_products or cat.destinations or cat.exporters):
            continue
        b.add_heading(cat.title.upper(), level=2)
        title = cat.title

        # -- Kenya's exports by product ------------------------------------
        if cat.kenya_products:
            rows = sorted(cat.kenya_products,
                          key=lambda r: r["years"].get(rev) or 0.0,
                          reverse=True)
            theme_rows = top_rows(rows, 10, years, "All other products")
            b._next_table(
                "Kenya's Exports of %s by Product, %d" % (title, rev), source)
            b.add_value_table(
                "Product", theme_rows, years, "Share in %d" % rev,
                "Kenya's Exports of %s by Product" % title, source,
                total_label="Total")
            _narrative_product(b, rows, title, years, rev,
                               cat.total_kenya(rev))
            pairs = _shares(theme_rows, years)
            if len(pairs) >= 2:
                img = make_donut(pairs, tmp_dir, "f_cat_prod.png",
                                 "Share of %s" % title)
                if img:
                    b._next_figure(
                        "Share of Kenya's Exports of %s by Product, %d"
                        % (title, rev), source)
                    b.add_figure(img)

        # -- Kenya's destination markets ------------------------------------
        if cat.destinations:
            rows = sorted(cat.destinations,
                          key=lambda r: r["years"].get(rev) or 0.0,
                          reverse=True)
            dest_rows = top_rows(rows, 12, years, "All other markets")
            b._next_table(
                "Destination Markets for Kenya's %s Exports, %d"
                % (title, rev), source)
            b.add_value_table(
                "Destination market", dest_rows, years, "Share in %d" % rev,
                "Kenya's %s Exports by Destination" % title, source,
                total_label="Total")
            _narrative_dest(b, rows, title, years, rev, cat.total_dest(rev))
            pairs = _shares(dest_rows, years)
            if len(pairs) >= 2:
                img = make_donut(pairs, tmp_dir, "f_cat_dest.png",
                                 "Kenya's %s Exports by Destination" % title)
                if img:
                    b._next_figure(
                        "Kenya's %s Exports by Destination, %d"
                        % (title, rev), source)
                    b.add_figure(img)

        # -- World exporters of the category -------------------------------
        if cat.exporters:
            rows = sorted(cat.exporters,
                          key=lambda r: r["years"].get(rev) or 0.0,
                          reverse=True)
            exp_rows = top_rows(rows, 12, years, "All other economies")
            b._next_table(
                "World Exports of %s by Economy, %d" % (title, rev), source)
            b.add_value_table(
                "Exporting economy", exp_rows, years, "Share in %d" % rev,
                "Countries Exporting %s" % title, source,
                total_label="Total")
            _narrative_exporters(b, rows, title, years, rev,
                                 cat.total_world(rev))


def section_whole(b, cfg, data, source, tmp_dir):
    """Commercial crafts as a whole: Kenya's portfolio, destinations and
    the global export picture (no import side is required)."""
    family = cfg.get("family_title", "Commercial Crafts")
    years = data.years
    rev = data.review_year

    b.add_heading("COMMERCIAL CRAFTS AS A WHOLE", level=1)

    # -- Kenya's exports by category --------------------------------------
    by_cat = data.kenya_by_category()
    if by_cat:
        b._next_table(
            "Trend on %s: Kenya's Exports by Category, %d" % (family, rev),
            source)
        b.add_value_table(
            "Category", by_cat, years, "Share in %d" % rev,
            "Kenya's Exports of %s by Category" % family, source,
            total_label="Total")
        totals = _year_totals(by_cat)
        total_last = totals.get(rev)
        if total_last:
            b.add_bullet("Kenya's total exports of %s were %s in %d."
                         % (family, usd_phrase(total_last), rev))
        lead_cat = by_cat[0]
        if lead_cat and total_last:
            b.add_bullet("The leading category was %s (%s; %.1f%% of Kenya's "
                         "craft exports)."
                         % (lead_cat["label"],
                            usd_phrase(lead_cat["years"].get(rev)),
                            (lead_cat["years"].get(rev) or 0.0)
                            / total_last * 100))
        follows = by_cat[1:3]
        if total_last and follows:
            b.add_bullet("It was followed by %s."
                         % ordinal_list(
                             ["%s (%s; %.1f%%)"
                              % (r["label"], usd_phrase(r["years"].get(rev)),
                                 (r["years"].get(rev) or 0.0)
                                 / total_last * 100)
                              for r in follows]))
        txt = _growth_sentence(by_cat, years, "Kenya's total exports of %s"
                               % family.lower())
        if txt:
            b.add_bullet(txt)
        pairs = _shares(by_cat, years)
        if len(pairs) >= 2:
            img = make_donut(pairs, tmp_dir, "f_whole_cat.png",
                             "Kenya's %s Exports by Category" % family)
            if img:
                b._next_figure("%s Exports by Category, %d"
                               % (family, rev), source)
                b.add_figure(img)

    # -- Kenya's craft exports by destination ------------------------------
    dest = data.destinations()
    if dest:
        dest_rows = top_rows(dest, 25, years, "All other markets")
        total_dest = sum((r["years"].get(rev) or 0.0) for r in dest)
        b._next_table(
            "Destination Markets for Kenya's %s Exports, %d"
            % (family, rev), source)
        b.add_value_table(
            "Destination market", dest_rows, years, "Share in %d" % rev,
            "Kenya's Exports of %s by Destination" % family, source,
            total_label="Total", rank=True)
        _narrative_dest(b, dest, family, years, rev, total_dest)
        pairs = _shares(dest_rows, years)
        if len(pairs) >= 2:
            img = make_donut(pairs, tmp_dir, "f_whole_dest.png",
                             "Kenya's %s Exports by Destination" % family)
            if img:
                b._next_figure(
                    "Kenya's %s Exports by Destination, %d" % (family, rev),
                    source)
                b.add_figure(img)

    # -- World exporters of commercial crafts ------------------------------
    exporters = data.exporters()
    if exporters:
        exp_rows = _ranked_rows(exporters, 12, years,
                                ensure_label="Kenya",
                                residual="All other economies")
        total_world = sum((r["years"].get(rev) or 0.0) for r in exporters)
        b._next_table(
            "World Exports of %s by Economy, %d" % (family, rev), source)
        b.add_value_table(
            "Exporting economy", exp_rows, years, "Share in %d" % rev,
            "Countries Exporting %s" % family, source,
            total_label="Total", rank=True)
        _narrative_exporters(b, exporters, family, years, rev, total_world)
        pairs = _shares(exp_rows, years)
        if len(pairs) >= 2:
            img = make_donut(pairs, tmp_dir, "f_whole_exp.png",
                             "Share of World Exports of %s" % family)
            if img:
                b._next_figure(
                    "World Exports of %s by Economy, %d" % (family, rev),
                    source)
                b.add_figure(img)

    # -- World exports of crafts by product --------------------------------
    wprod = data.world_products()
    if wprod:
        wprod_rows = top_rows(wprod, 15, years, "All other products")
        total_wprod = sum((r["years"].get(rev) or 0.0) for r in wprod)
        b._next_table(
            "Trend on %s Globally - Export, %d" % (family, rev), source)
        b.add_value_table(
            "Product", wprod_rows, years, "Share in %d" % rev,
            "Global Exports of %s by Product" % family, source,
            total_label="Total")
        _narrative_world_products(b, wprod, years, rev, total_wprod)

    # -- Export specialization (share of Kenya's total exports) ------------
    merch = data.kenya_merchandise_total()
    kenya_crafts = sum((r["years"].get(rev) or 0.0)
                       for r in by_cat) if by_cat else 0.0
    if merch.get(rev) and kenya_crafts:
        b.add_bullet(
            "%s accounted for %.2f%% of Kenya's total merchandise exports "
            "in %d (%s of %s), reflecting Kenya's export specialization in "
            "the crafts portfolio."
            % (family, kenya_crafts / merch.get(rev) * 100, rev,
               usd_phrase(kenya_crafts), usd_phrase(merch.get(rev))))


def build_crafts_report(cfg, data_dir, out_path, tmp_dir):
    """Build the crafts Word report to ``out_path``."""
    os.makedirs(tmp_dir, exist_ok=True)
    source = cfg.get("source", "International Trade Centre (ITC), Trade Map")
    data = CraftsData(data_dir)
    b = ProfileBuilder(cfg, {})
    b.title_page(cfg)
    family = cfg.get("family_title", "Commercial Crafts")
    b.add_heading("TRADE IN %s" % family.upper(), level=1)
    for par in cfg.get("intro", []):
        b.add_para(par)
    section_categories(b, cfg, data, source, tmp_dir)
    section_whole(b, cfg, data, source, tmp_dir)
    return b.doc


# --------------------------------------------------------------------------
# Excel deliverable
# --------------------------------------------------------------------------
def write_crafts_excel(cfg, data_dir, out_path):
    """Companion workbook with the same tables."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    NAVY = "1F3864"
    hdr_fill = PatternFill("solid", fgColor=NAVY)
    cm = Alignment(horizontal="center")
    lm = Alignment(horizontal="left")
    data = CraftsData(data_dir)
    years = data.years
    rev = data.review_year
    wb = Workbook()
    wb.remove(wb.active)

    def value_sheet(title, first_col, rows, label_key="label",
                    code_key=None):
        ws = wb.create_sheet(title[:31])
        hdr = [first_col] + list(years) + ["Share in %d" % rev]
        for c, h in enumerate(hdr, 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = hdr_fill
            cell.alignment = cm
        rev_total = sum(display(r["years"].get(rev)) or 0.0 for r in rows)
        for ri, r in enumerate(rows, start=2):
            lab = str(r.get(label_key) or r.get("label") or "")
            if code_key and r.get(code_key):
                lab = "%s  %s" % (r.get(code_key), short_label(lab, 44))
            ws.cell(ri, 1, lab).alignment = lm
            for i, y in enumerate(years, start=2):
                v = display(r["years"].get(y))
                ws.cell(ri, i, None if v is None else round(v, 1))
                ws.cell(ri, i).alignment = cm
            v = display(r["years"].get(rev))
            share = (v / rev_total if v is not None and rev_total else None)
            ws.cell(ri, len(years) + 2, share).number_format = "0.0%"
        ws.column_dimensions["A"].width = min(
            60, max(30, *(len(str(r.get("label"))) for r in rows)))
        return ws

    by_cat = data.kenya_by_category()
    if by_cat:
        value_sheet("Kenya Exports by Category", "Category", by_cat)
    dest = top_rows(data.destinations(), 25, years, "All other markets")
    if dest:
        value_sheet("Kenya Exports by Destination", "Destination", dest)
    exp = _ranked_rows(data.exporters(), 12, years, ensure_label="Kenya",
                       residual="All other economies")
    if exp:
        value_sheet("World Exporters", "Exporting economy", exp)
    wprod = top_rows(data.world_products(), 15, years, "All other products")
    if wprod:
        value_sheet("Global Exports by Product", "Product", wprod,
                    label_key="label", code_key="code")
    for cat in data.order:
        rows = sorted(cat.kenya_products,
                      key=lambda r: r["years"].get(rev) or 0.0,
                      reverse=True)
        if rows:
            value_sheet("Products " + cat.title, "Product", rows,
                        label_key="label", code_key="code")
        rows = sorted(cat.destinations,
                      key=lambda r: r["years"].get(rev) or 0.0,
                      reverse=True)
        if rows:
            value_sheet("Destinations " + cat.title, "Destination", rows)
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    return wb


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(
        description="Generate a product-profile report for Kenya's crafts.")
    ap.add_argument("--config", default=None)
    ap.add_argument("--data-dir", default=os.path.join(BASE_DIR, "crafts"))
    ap.add_argument(
        "--output",
        default=os.path.join(BASE_DIR, "output",
                             "COMMERCIAL CRAFTS PRODUCT PROFILE.docx"))
    ap.add_argument("--tmp",
                    default=os.path.join(BASE_DIR, "output", ".tmp"))
    args = ap.parse_args()

    import json
    if args.config:
        cfg = json.load(open(args.config, encoding="utf-8"))
    else:
        cfg_path = os.path.join(BASE_DIR, "config",
                                "product_profile_crafts.json")
        cfg = json.load(open(cfg_path, encoding="utf-8"))

    data = CraftsData(args.data_dir)
    print("[1/4] Loading craft product-group files from: %s" % args.data_dir)
    print("      period    = %s - %s"
          % (data.start_year, data.review_year))
    print("      categories= %d" % len(data.categories))
    for c in data.order:
        print("        - %-48s %s"
              % (c.title, usd_phrase(c.total_kenya(data.review_year))))
    for w in data.warnings:
        print("      [warn] %s" % w)
    if not data.categories:
        print("[ERROR] No craft files could be loaded from %r" % args.data_dir)
        return 1

    print("[2/4] Building report             : %s" % args.output)
    doc = build_crafts_report(cfg, args.data_dir, args.output, args.tmp)
    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
    doc.save(args.output)
    print("[3/4] Report saved to             : %s" % args.output)
    t = os.path.splitext(args.output)[0] + " TABLES.xlsx"
    try:
        write_crafts_excel(cfg, args.data_dir, t)
        print("[4/4] Excel deliverable saved to : %s" % t)
    except Exception as e:
        print("      [warn] Excel deliverable skipped: %s" % e)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())